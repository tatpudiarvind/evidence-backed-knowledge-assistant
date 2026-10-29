"""Helpers shared by the tests: settings, sample files, API calls, and a stand-in embedder."""

import hashlib
import io
import math
import re
import zipfile
from collections.abc import Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import docx
from fastapi import FastAPI
from fastapi.testclient import TestClient
from httpx2 import Response
from pypdf import PdfReader, PdfWriter

from app.core.config import Settings
from app.domain.indexing import EMBEDDING_DIMENSION
from app.services.ingestion import DocumentIndexer

CSRF_HEADERS = {"X-CSRF-Protection": "1"}
# Port 9 (discard) refuses connections, so unit tests fail fast if they touch the database.
UNREACHABLE_DATABASE_URL = "postgresql+psycopg://nobody:nothing@127.0.0.1:9/unit_tests"
DEFAULT_PASSWORD = "correct horse battery staple"
TXT_MIME = "text/plain"
PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def make_settings(**overrides: Any) -> Settings:
    values: dict[str, Any] = {
        "database_url": UNREACHABLE_DATABASE_URL,
        "database_connect_timeout_seconds": 1,
        # Uploads stay queued; tests that need indexing call index_queued_documents.
        "ingestion_worker_enabled": False,
    }
    values.update(overrides)
    return Settings(_env_file=None, **values)


def pdf_bytes(*pages: str) -> bytes:
    """A real PDF with one page per argument, showing its lines of text in Helvetica.

    An empty string makes a page without any text. Text must fit the Windows-1252 code page.
    """
    font = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>"
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"", font]
    page_numbers: list[int] = []
    for text in pages or ("Synthetic test PDF",):
        shown = b" T* ".join(_pdf_string(line) + b" Tj" for line in text.split("\n"))
        stream = b"BT /F1 12 Tf 14 TL 72 720 Td " + shown + b" ET" if text else b""
        objects.append(b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream))
        objects.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 3 0 R >> >> /Contents %d 0 R >>" % len(objects)
        )
        page_numbers.append(len(objects))
    kids = b" ".join(b"%d 0 R" % number for number in page_numbers)
    objects[1] = b"<< /Type /Pages /Kids [%s] /Count %d >>" % (kids, len(page_numbers))
    output = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(output))
        output += b"%d 0 obj\n%s\nendobj\n" % (number, body)
    xref_offset = len(output)
    output += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    output += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    output += b"trailer\n<< /Size %d /Root 1 0 R >>\n" % (len(objects) + 1)
    output += b"startxref\n%d\n%%%%EOF\n" % xref_offset
    return bytes(output)


def _pdf_string(text: str) -> bytes:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    return b"(" + escaped.encode("cp1252") + b")"


def encrypted_pdf_bytes(text: str, *, user_password: str) -> bytes:
    """A PDF encrypted with AES-256. With an empty user password, anyone can still open it."""
    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(pdf_bytes(text))))
    writer.encrypt(user_password, "synthetic owner password", algorithm="AES-256")
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


@dataclass(frozen=True)
class Heading:
    text: str
    level: int = 1


def docx_bytes(*content: str | Heading | list[list[str]]) -> bytes:
    """A real Word document: strings are paragraphs, `Heading`s headings, lists of rows tables."""
    document = docx.Document()
    for item in content or ("Synthetic test document",):
        if isinstance(item, Heading):
            document.add_heading(item.text, item.level)
        elif isinstance(item, str):
            document.add_paragraph(item)
        else:
            table = document.add_table(rows=len(item), cols=len(item[0]))
            for row, values in zip(table.rows, item, strict=True):
                for cell, value in zip(row.cells, values, strict=True):
                    cell.text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def zip_bytes(entries: dict[str, str]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for name, content in entries.items():
            archive.writestr(name, content)
    return buffer.getvalue()


def register(
    client: TestClient, email: str = "reader@example.com", password: str = DEFAULT_PASSWORD
) -> Response:
    return client.post(
        "/api/v1/auth/register",
        json={"email": email, "password": password},
        headers=CSRF_HEADERS,
    )


def login(
    client: TestClient, email: str = "reader@example.com", password: str = DEFAULT_PASSWORD
) -> Response:
    return client.post(
        "/api/v1/auth/login",
        json={"email": email, "password": password},
        headers=CSRF_HEADERS,
    )


def create_workspace(client: TestClient, name: str = "Research notes") -> dict[str, Any]:
    response = client.post("/api/v1/workspaces", json={"name": name}, headers=CSRF_HEADERS)
    assert response.status_code == 201, response.text
    workspace: dict[str, Any] = response.json()
    return workspace


def upload(
    client: TestClient,
    workspace_id: str,
    filename: str,
    content: bytes,
    mime_type: str = "application/octet-stream",
) -> Response:
    return client.post(
        f"/api/v1/workspaces/{workspace_id}/documents",
        files={"file": (filename, content, mime_type)},
        headers=CSRF_HEADERS,
    )


def error_code(response: Response) -> str:
    code: str = response.json()["error"]["code"]
    return code


def stored_files(storage_dir: Path) -> list[Path]:
    """Every file under the storage folder, including leftover temporary uploads."""
    return sorted(path for path in storage_dir.rglob("*") if path.is_file())


class HashingEmbeddings:
    """A fast, repeatable stand-in for the embedding model, so tests need no download.

    Each word adds 1 or -1 at a position picked by hashing it, so texts sharing words get
    similar vectors. Records the texts it was given.
    """

    def __init__(self) -> None:
        self.document_batches: list[list[str]] = []

    def embed_documents(self, texts: Sequence[str]) -> list[list[float]]:
        self.document_batches.append(list(texts))
        return [hashed_vector(text) for text in texts]

    def embed_query(self, text: str) -> list[float]:
        return hashed_vector(text)


def hashed_vector(text: str) -> list[float]:
    vector = [0.0] * EMBEDDING_DIMENSION
    for word in re.findall(r"\w+", text.lower()):
        digest = hashlib.blake2b(word.encode(), digest_size=8).digest()
        position = int.from_bytes(digest[:4], "little") % EMBEDDING_DIMENSION
        vector[position] += 1.0 if digest[4] & 1 else -1.0
    norm = math.sqrt(sum(value * value for value in vector))
    if norm == 0:
        # No words, or words that cancel out: any fixed unit vector will do.
        vector[0], norm = 1.0, 1.0
    return [value / norm for value in vector]


def index_queued_documents(app: FastAPI) -> int:
    """Index every queued document now, as the background worker would. Returns how many."""
    state = app.state
    indexer = DocumentIndexer(
        state.session_factory, state.storage, state.embeddings, state.settings
    )
    count = 0
    while indexer.process_next():
        count += 1
    return count
