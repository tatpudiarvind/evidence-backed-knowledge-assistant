"""Text extraction from PDF, Word, and text files (P3-02) within limits (P3-10). No database."""

import io
from collections.abc import Callable
from typing import Any

import docx
import pytest

from app.domain.indexing import IndexingError, IndexingFailure, TextBlock
from app.providers import parsers
from app.providers.parsers import (
    DocxParser,
    ParseLimits,
    Parser,
    PdfParser,
    TextParser,
    parser_for,
)
from tests.support import (
    DOCX_MIME,
    PDF_MIME,
    TXT_MIME,
    Heading,
    docx_bytes,
    encrypted_pdf_bytes,
    pdf_bytes,
    zip_bytes,
)

DEFAULT_LIMITS: dict[str, Any] = {"max_pages": 10, "max_characters": 10_000, "timeout_seconds": 60}
MakeParser = Callable[[ParseLimits], Parser]


def _extract(parser: MakeParser, content: bytes, **limits: Any) -> list[TextBlock]:
    return parser(ParseLimits(**{**DEFAULT_LIMITS, **limits})).extract(io.BytesIO(content))


def _failure(parser: MakeParser, content: bytes, **limits: Any) -> IndexingFailure:
    with pytest.raises(IndexingError) as raised:
        _extract(parser, content, **limits)
    return raised.value.failure


def test_pdf_pages_become_numbered_blocks() -> None:
    content = pdf_bytes("Retention policy\nKeep records (synthetic) for 7 years.", "", "Appendix")

    assert _extract(PdfParser, content) == [
        TextBlock("Retention policy\nKeep records (synthetic) for 7 years.", page_number=1),
        TextBlock("", page_number=2),
        TextBlock("Appendix", page_number=3),
    ]


def test_pdfs_encrypted_without_a_reading_password_are_read() -> None:
    content = encrypted_pdf_bytes("Open to read", user_password="")

    assert _extract(PdfParser, content) == [TextBlock("Open to read", page_number=1)]


def test_password_protected_pdfs_are_refused() -> None:
    content = encrypted_pdf_bytes("Secret", user_password="synthetic password")

    assert _failure(PdfParser, content) == IndexingFailure.DOCUMENT_ENCRYPTED


def test_pdf_page_limit() -> None:
    content = pdf_bytes("one", "two", "three")

    assert len(_extract(PdfParser, content, max_pages=3)) == 3
    assert _failure(PdfParser, content, max_pages=2) == IndexingFailure.TOO_MANY_PAGES


def test_pdf_text_limit() -> None:
    content = pdf_bytes("x" * 600, "y" * 600)

    assert _failure(PdfParser, content, max_characters=1000) == IndexingFailure.TOO_MUCH_TEXT


def test_pdf_time_limit() -> None:
    assert _failure(PdfParser, pdf_bytes(), timeout_seconds=0) == IndexingFailure.PARSE_TIMEOUT


@pytest.mark.parametrize(
    "content",
    [
        b"%PDF-1.7\nnot really a PDF",
        pdf_bytes("Cut short")[:200],
        b"%PDF-1.4\n1 0 obj\n<< >>\nendobj\ntrailer\n<< >>\n%%EOF\n",
    ],
    ids=["garbage", "truncated", "no-pages"],
)
def test_damaged_pdfs_are_unreadable(content: bytes) -> None:
    assert _failure(PdfParser, content) == IndexingFailure.DOCUMENT_UNREADABLE


def test_word_sections_start_at_headings() -> None:
    content = docx_bytes(
        "Before any heading.",
        Heading("Background"),
        "First paragraph.",
        "Second paragraph.",
        Heading("Results", 2),
        "Findings.",
    )

    assert _extract(DocxParser, content) == [
        TextBlock("Before any heading."),
        TextBlock("Background\n\nFirst paragraph.\n\nSecond paragraph.", section="Background"),
        TextBlock("Results\n\nFindings.", section="Results"),
    ]


def test_titles_count_as_headings_and_empty_headings_do_not() -> None:
    content = docx_bytes(Heading("Report", 0), "One.", Heading("", 2), "Two.")

    assert _extract(DocxParser, content) == [
        TextBlock(""),
        TextBlock("Report\n\nOne.\n\n\n\nTwo.", section="Report"),
    ]


def test_word_tables_become_rows_of_cells() -> None:
    content = docx_bytes("Intro.", [["Name", "Value"], ["Alpha", "1"], ["", ""]])

    assert _extract(DocxParser, content) == [TextBlock("Intro.\n\nName | Value\nAlpha | 1")]


def test_merged_and_nested_table_cells_appear_once() -> None:
    document = docx.Document()
    table = document.add_table(rows=2, cols=3)
    merged = table.cell(0, 0).merge(table.cell(0, 1))
    merged.text = "Merged across"
    table.cell(0, 2).text = "Last"
    table.cell(1, 0).text = "Outer"
    nested = table.cell(1, 1).add_table(rows=1, cols=2)
    nested.cell(0, 0).text = "Inner a"
    nested.cell(0, 1).text = "Inner b"
    buffer = io.BytesIO()
    document.save(buffer)

    assert _extract(DocxParser, buffer.getvalue()) == [
        TextBlock("Merged across | Last\nOuter | Inner a | Inner b")
    ]


def test_word_text_limit() -> None:
    content = docx_bytes("x" * 600, "y" * 600)

    assert _failure(DocxParser, content, max_characters=1000) == IndexingFailure.TOO_MUCH_TEXT


def test_word_time_limit() -> None:
    assert _failure(DocxParser, docx_bytes(), timeout_seconds=0) == IndexingFailure.PARSE_TIMEOUT


def test_word_files_with_too_much_xml_are_refused(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(parsers, "DOCX_MAX_XML_BYTES", 1000)

    assert _failure(DocxParser, docx_bytes()) == IndexingFailure.TOO_MUCH_TEXT


@pytest.mark.parametrize(
    "content",
    [
        zip_bytes({"[Content_Types].xml": "<Types/>", "word/document.xml": "<document/>"}),
        docx_bytes()[:500],
        b"PK\x03\x04 not a real archive",
    ],
    ids=["no-relationships", "truncated", "garbage"],
)
def test_damaged_word_files_are_unreadable(content: bytes) -> None:
    assert _failure(DocxParser, content) == IndexingFailure.DOCUMENT_UNREADABLE


def test_text_files_are_one_block_without_a_byte_order_mark() -> None:
    content = "\ufeffLine one\nLine two, caf\u00e9".encode()

    assert _extract(TextParser, content) == [TextBlock("Line one\nLine two, caf\u00e9")]


def test_text_character_limit() -> None:
    at_limit = "\u00e9" * 1000

    assert _extract(TextParser, at_limit.encode(), max_characters=1000) == [TextBlock(at_limit)]
    over = (at_limit + "x").encode()
    assert _failure(TextParser, over, max_characters=1000) == IndexingFailure.TOO_MUCH_TEXT


def test_text_files_are_measured_before_decoding() -> None:
    # Four bytes per character is UTF-8's most, so 4,000 bytes can still be 1,000 characters.
    emoji = "\U0001f600" * 1000

    assert _extract(TextParser, emoji.encode(), max_characters=1000) == [TextBlock(emoji)]
    too_big = b"x" * 4001
    assert _failure(TextParser, too_big, max_characters=1000) == IndexingFailure.TOO_MUCH_TEXT


def test_text_files_that_are_not_utf8_are_unreadable() -> None:
    assert _failure(TextParser, b"caf\xe9") == IndexingFailure.DOCUMENT_UNREADABLE


@pytest.mark.parametrize(
    ("mime_type", "parser"),
    [(PDF_MIME, PdfParser), (DOCX_MIME, DocxParser), (TXT_MIME, TextParser)],
)
def test_each_stored_type_has_a_parser(mime_type: str, parser: type) -> None:
    assert isinstance(parser_for(mime_type, ParseLimits(**DEFAULT_LIMITS)), parser)


def test_other_types_are_unreadable() -> None:
    with pytest.raises(IndexingError) as raised:
        parser_for("application/zip", ParseLimits(**DEFAULT_LIMITS))

    assert raised.value.failure == IndexingFailure.DOCUMENT_UNREADABLE
